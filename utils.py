import json
import os
from datetime import datetime
from io import BytesIO
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Preformatted
from xml.sax.saxutils import escape
import re

HISTORY_FILE = "history.json"


# ========== Chat Messages =========
def add_message(history, role, message):
    if not message or not message.strip():
        return

    history.append(
        {
            "role": role,
            "message": message,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    )

# =========== history ===========
def load_history():
    if not os.path.exists(HISTORY_FILE):
        return []

    try:
        with open(HISTORY_FILE, "r", encoding="utf-8") as f:
            return json.load(f)

    except Exception:
        return []


def save_history(history):
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(
            history,
            f,
            indent=4,
            ensure_ascii=False,
        )


# ========== export helpers ==========
# FIX: exports used to dump the raw internal tags straight into the
# file - e.g. "[IMAGE]/tmp/tmpabc123.png" or
# "[DOCUMENT]report.pdf::/home/.../uploads/173..._report.pdf" showed up
# as literal text in the exported PDF/DOCX/TXT/MD. This converts those
# internal markers into readable lines before export.
def _humanize_message(text):
    if text is None:
        return ""

    if "[IMAGE]" in text:
        path = text.replace("[IMAGE]", "").strip()
        return f"📷 Image attached: {os.path.basename(path)}"

    if "[VIDEO]" in text:
        path = text.replace("[VIDEO]", "").strip()
        return f"🎬 Video attached: {os.path.basename(path)}"

    if "[DOCUMENT]" in text:
        payload = text.replace("[DOCUMENT]", "").strip()
        name = payload.split("::", 1)[0] if "::" in payload else payload
        return f"📄 Document attached: {name}"

    if text.startswith("[URL]"):
        link = text.replace("[URL]", "").strip()
        return f"🔗 Link: {link}"

    return text


# ========= export =============
def export_json(messages):
    data = []

    for msg in messages:
        data.append({
            "id": msg["id"],
            "role": msg["role"],
            "message": msg["message"],
            "created_at": str(msg["created_at"])
        })

    return json.dumps(data, indent=4, ensure_ascii=False).encode("utf-8")


def export_txt(messages):
    text = "\n\n".join(
        f"{m['role'].upper()}:\n{_humanize_message(m['message'])}"
        for m in messages
    )
    return text.encode("utf-8")


def export_md(messages):
    md = "\n\n".join(
        f"## {m['role'].capitalize()}\n\n{_humanize_message(m['message'])}"
        for m in messages
    )
    return md.encode("utf-8")


def export_docx(messages):
    doc = Document()

    doc.add_heading("Chat History", level=1)

    for m in messages:
        doc.add_heading(m["role"].capitalize(), level=2)
        doc.add_paragraph(_humanize_message(m["message"]))

    file = BytesIO()
    doc.save(file)
    file.seek(0)

    return file


def export_pdf(messages):
    file = BytesIO()

    doc = SimpleDocTemplate(file)
    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph("<b>Chat History</b>", styles["Title"])
    )

    for m in messages:

        story.append(
            Paragraph(
                f"<b>{escape(m['role'].capitalize())}</b>",
                styles["Heading2"]
            )
        )

        text = _humanize_message(m["message"])

        # Check for fenced code blocks
        if "```" in text:
            clean = re.sub(r"```[a-zA-Z]*", "", text)
            clean = clean.replace("```", "")
            story.append(
                Preformatted(clean, styles["Code"])
            )
        else:
            safe = escape(text).replace("\n", "<br/>")
            story.append(
                Paragraph(safe, styles["BodyText"])
            )

    doc.build(story)

    file.seek(0)
    return file

# ========== analytics ==============
def conversation_stats(history):
    total_messages = len(history)
    user_messages = sum(
        1 for msg in history
        if msg["role"] == "user"
    )

    assistant_messages = sum(
        1 for msg in history
        if msg["role"] == "assistant"
    )

    total_words = sum(
        len(msg["message"].split())
        for msg in history
    )

    if history:
        start = datetime.strptime(
            history[0]["timestamp"],
            "%Y-%m-%d %H:%M:%S",
        )

        end = datetime.strptime(
            history[-1]["timestamp"],
            "%Y-%m-%d %H:%M:%S",
        )

        duration = end - start

    else:
        duration = "0:00:00"

    return {
        "Total Messages": total_messages,
        "User Messages": user_messages,
        "AI Messages": assistant_messages,
        "Total Words": total_words,
        "Conversation Duration": str(duration),
    }


def get_last_response(history):
    for msg in reversed(history):
        if msg["role"] == "assistant":
            return msg["message"]
    return ""